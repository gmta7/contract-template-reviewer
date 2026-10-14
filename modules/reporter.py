"""
审核报告生成模块

支持两种输出格式：
- HTML 报告：供法务在浏览器中直观审阅
- Word 报告：供法务向业务方反馈（.docx）
"""
import html as html_lib
import os
from datetime import datetime
from typing import Optional

from modules.differ import DiffResult, DiffType


# ─────────────────────────────────────────────
# HTML 报告
# ─────────────────────────────────────────────
_HTML_TEMPLATE = """\
<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1.0"/>
<title>合同审核报告 - {contract_name}</title>
<style>
  :root {{
    --green: #16a34a; --green-bg: #f0fdf4; --green-border: #bbf7d0;
    --red: #dc2626;   --red-bg: #fef2f2;   --red-border: #fecaca;
    --yellow: #d97706; --yellow-bg: #fffbeb; --yellow-border: #fde68a;
    --gray: #6b7280;  --gray-bg: #f9fafb;   --gray-border: #e5e7eb;
    --blue: #2563eb;  --blue-bg: #eff6ff;   --blue-border: #bfdbfe;
  }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: "Microsoft YaHei", "PingFang SC", sans-serif; font-size: 14px; color: #111827; background: #f3f4f6; }}
  .page {{ max-width: 960px; margin: 0 auto; padding: 24px 16px 48px; }}
  h1 {{ font-size: 22px; font-weight: 700; margin-bottom: 4px; }}
  .meta {{ color: var(--gray); font-size: 13px; margin-bottom: 24px; }}

  /* 汇总卡片 */
  .summary-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(160px,1fr)); gap: 12px; margin-bottom: 24px; }}
  .card {{ background: #fff; border-radius: 8px; padding: 16px; border: 1px solid var(--gray-border); }}
  .card .num {{ font-size: 28px; font-weight: 700; }}
  .card .label {{ font-size: 12px; color: var(--gray); margin-top: 4px; }}
  .card.green .num {{ color: var(--green); }}
  .card.red .num {{ color: var(--red); }}
  .card.yellow .num {{ color: var(--yellow); }}
  .card.blue .num {{ color: var(--blue); }}

  /* 告警横幅 */
  .alert-banner {{ background: var(--red-bg); border: 1px solid var(--red-border); border-radius: 8px; padding: 12px 16px; margin-bottom: 24px; color: var(--red); font-weight: 600; display: flex; align-items: center; gap: 8px; }}
  .ok-banner {{ background: var(--green-bg); border: 1px solid var(--green-border); border-radius: 8px; padding: 12px 16px; margin-bottom: 24px; color: var(--green); font-weight: 600; display: flex; align-items: center; gap: 8px; }}

  /* 区块 */
  .section {{ background: #fff; border-radius: 8px; border: 1px solid var(--gray-border); margin-bottom: 16px; overflow: hidden; }}
  .section-header {{ padding: 12px 16px; font-weight: 600; font-size: 14px; border-bottom: 1px solid var(--gray-border); display: flex; align-items: center; gap: 8px; cursor: pointer; user-select: none; }}
  .section-header .badge {{ font-size: 12px; padding: 2px 8px; border-radius: 9999px; font-weight: 500; }}
  .badge-red {{ background: var(--red-bg); color: var(--red); }}
  .badge-green {{ background: var(--green-bg); color: var(--green); }}
  .badge-yellow {{ background: var(--yellow-bg); color: var(--yellow); }}
  .section-body {{ padding: 12px 16px; }}

  /* 差异行 */
  .diff-row {{ border-radius: 6px; padding: 10px 12px; margin-bottom: 8px; font-size: 13px; line-height: 1.8; }}
  .diff-row.filled {{ background: var(--green-bg); border-left: 3px solid var(--green); }}
  .diff-row.tampered {{ background: var(--red-bg); border-left: 3px solid var(--red); }}
  .diff-row.added {{ background: var(--blue-bg); border-left: 3px solid var(--blue); }}
  .diff-row.removed {{ background: var(--yellow-bg); border-left: 3px solid var(--yellow); }}
  .diff-label {{ font-size: 11px; font-weight: 600; text-transform: uppercase; letter-spacing: .5px; margin-bottom: 6px; }}
  .diff-row.filled .diff-label {{ color: var(--green); }}
  .diff-row.tampered .diff-label {{ color: var(--red); }}
  .diff-row.added .diff-label {{ color: var(--blue); }}
  .diff-row.removed .diff-label {{ color: var(--yellow); }}
  .diff-cols {{ display: grid; grid-template-columns: 1fr 1fr; gap: 8px; }}
  .diff-col-header {{ font-size: 11px; color: var(--gray); margin-bottom: 4px; }}
  .diff-text {{ background: rgba(0,0,0,.04); border-radius: 4px; padding: 6px 8px; word-break: break-all; }}
  span.diff-del {{ background: #fca5a5; border-radius: 2px; padding: 0 1px; text-decoration: line-through; }}
  span.diff-ins {{ background: #86efac; border-radius: 2px; padding: 0 1px; }}

  /* LLM 结果 */
  .llm-section {{ background: var(--blue-bg); border: 1px solid var(--blue-border); border-radius: 8px; padding: 16px; margin-bottom: 16px; }}
  .llm-section h3 {{ color: var(--blue); margin-bottom: 12px; }}
  table {{ width: 100%; border-collapse: collapse; margin-top: 8px; }}
  th {{ background: var(--gray-bg); text-align: left; padding: 6px 10px; font-size: 12px; border-bottom: 2px solid var(--gray-border); }}
  td {{ padding: 8px 10px; border-bottom: 1px solid var(--gray-border); font-size: 13px; vertical-align: top; }}
  .severity-high {{ color: var(--red); font-weight: 600; }}
  .severity-medium {{ color: var(--yellow); font-weight: 600; }}
  .severity-low {{ color: var(--green); font-weight: 600; }}

  @media print {{
    body {{ background: white; }}
    .page {{ padding: 0; }}
    .no-print {{ display: none; }}
  }}
</style>
</head>
<body>
<div class="page">
  <h1>合同审核报告</h1>
  <div class="meta">模版：{template_name} &nbsp;|&nbsp; 回签合同：{contract_name} &nbsp;|&nbsp; 审核时间：{timestamp} &nbsp;|&nbsp; 相似度：{similarity:.1%}</div>

  {banner}

  <div class="summary-grid">
    <div class="card green"><div class="num">{filled}</div><div class="label">正常填写字段</div></div>
    <div class="card red"><div class="num">{tampered}</div><div class="label">疑似修改条款</div></div>
    <div class="card yellow"><div class="num">{removed}</div><div class="label">模版内容缺失</div></div>
    <div class="card blue"><div class="num">{added}</div><div class="label">新增内容</div></div>
  </div>

  {sections_html}

  {llm_html}

</div>
<script>
document.querySelectorAll('.section-header').forEach(h=>{{
  h.addEventListener('click',()=>{{
    const body = h.nextElementSibling;
    if(body) body.style.display = body.style.display==='none' ? '' : 'none';
  }});
}});
</script>
</body>
</html>
"""


def _escape(text: str) -> str:
    return html_lib.escape(str(text)) if text else ""


def _build_sections_html(diff_result: DiffResult) -> str:
    """生成各差异类型的 HTML 区块"""
    filled_rows = []
    tampered_rows = []
    added_rows = []
    removed_rows = []

    for item in diff_result.items:
        if item.type == DiffType.FILLED:
            filled_rows.append(
                f'<div class="diff-row filled">'
                f'<div class="diff-label">✅ 填写字段</div>'
                f'<div class="diff-cols">'
                f'<div><div class="diff-col-header">模版原文（含空白）</div>'
                f'<div class="diff-text">{_escape(item.template_text)}</div></div>'
                f'<div><div class="diff-col-header">回签填写内容</div>'
                f'<div class="diff-text">{_escape(item.contract_text)}</div></div>'
                f'</div></div>'
            )
        elif item.type == DiffType.TAMPERED:
            tampered_rows.append(
                f'<div class="diff-row tampered">'
                f'<div class="diff-label">🚨 疑似修改</div>'
                f'<div class="diff-cols">'
                f'<div><div class="diff-col-header">模版原文</div>'
                f'<div class="diff-text">{item.template_html or _escape(item.template_text)}</div></div>'
                f'<div><div class="diff-col-header">回签版本（修改处已标红/绿）</div>'
                f'<div class="diff-text">{item.contract_html or _escape(item.contract_text)}</div></div>'
                f'</div></div>'
            )
        elif item.type == DiffType.ADDED:
            added_rows.append(
                f'<div class="diff-row added">'
                f'<div class="diff-label">➕ 新增内容（模版外）</div>'
                f'<div class="diff-text">{_escape(item.contract_text)}</div>'
                f'</div>'
            )
        elif item.type == DiffType.REMOVED:
            removed_rows.append(
                f'<div class="diff-row removed">'
                f'<div class="diff-label">⚠️ 内容缺失</div>'
                f'<div class="diff-text">{_escape(item.template_text)}</div>'
                f'</div>'
            )

    parts = []

    if tampered_rows:
        parts.append(
            f'<div class="section">'
            f'<div class="section-header">🚨 疑似修改条款 <span class="badge badge-red">{len(tampered_rows)} 处</span></div>'
            f'<div class="section-body">{"".join(tampered_rows)}</div>'
            f'</div>'
        )

    if removed_rows:
        parts.append(
            f'<div class="section">'
            f'<div class="section-header">⚠️ 模版内容缺失 <span class="badge badge-yellow">{len(removed_rows)} 处</span></div>'
            f'<div class="section-body">{"".join(removed_rows)}</div>'
            f'</div>'
        )

    if added_rows:
        parts.append(
            f'<div class="section">'
            f'<div class="section-header">➕ 模版外新增内容 <span class="badge badge-red">{len(added_rows)} 处</span></div>'
            f'<div class="section-body">{"".join(added_rows)}</div>'
            f'</div>'
        )

    if filled_rows:
        body_hidden = 'style="display:none"' if len(filled_rows) > 20 else ''
        parts.append(
            f'<div class="section">'
            f'<div class="section-header">✅ 正常填写字段 <span class="badge badge-green">{len(filled_rows)} 处</span>（点击展开）</div>'
            f'<div class="section-body" {body_hidden}>{"".join(filled_rows)}</div>'
            f'</div>'
        )

    return "".join(parts)


def _build_llm_html(llm_result: Optional[dict]) -> str:
    if not llm_result or "error" in llm_result:
        return ""

    parts = ['<div class="llm-section"><h3>🤖 LLM 辅助分析结果</h3>']

    summary = llm_result.get("summary", "")
    risk = llm_result.get("risk_level", "")
    rec = llm_result.get("recommendation", "")
    risk_cls = f"severity-{risk}" if risk in ("high", "medium", "low") else ""
    risk_label = {"high": "高风险", "medium": "中等风险", "low": "低风险"}.get(risk, risk)

    if summary:
        parts.append(f'<p><strong>摘要：</strong>{_escape(summary)}</p>')
    if risk:
        parts.append(f'<p style="margin-top:8px"><strong>风险等级：</strong><span class="{risk_cls}">{risk_label}</span></p>')
    if rec:
        parts.append(f'<p style="margin-top:4px"><strong>建议：</strong>{_escape(rec)}</p>')

    filled = llm_result.get("filled_fields", [])
    if filled:
        parts.append('<h4 style="margin:16px 0 8px">填写字段清单</h4><table><tr><th>字段</th><th>填写值</th><th>位置</th></tr>')
        for f in filled:
            parts.append(
                f'<tr><td>{_escape(f.get("field_name",""))}</td>'
                f'<td>{_escape(f.get("value",""))}</td>'
                f'<td>{_escape(f.get("location",""))}</td></tr>'
            )
        parts.append("</table>")

    tampered = llm_result.get("tampered_sections", [])
    if tampered:
        parts.append('<h4 style="margin:16px 0 8px">修改条款明细</h4><table><tr><th>模版原文</th><th>修改后</th><th>位置</th><th>严重程度</th><th>风险说明</th></tr>')
        for t in tampered:
            sev = t.get("severity", "")
            sev_cls = f"severity-{sev}" if sev in ("high", "medium", "low") else ""
            sev_label = {"high": "高", "medium": "中", "low": "低"}.get(sev, sev)
            parts.append(
                f'<tr><td>{_escape(t.get("original",""))}</td>'
                f'<td>{_escape(t.get("modified",""))}</td>'
                f'<td>{_escape(t.get("location",""))}</td>'
                f'<td class="{sev_cls}">{sev_label}</td>'
                f'<td>{_escape(t.get("reason",""))}</td></tr>'
            )
        parts.append("</table>")

    parts.append("</div>")
    return "".join(parts)


def generate_html_report(
    diff_result: DiffResult,
    template_name: str,
    contract_name: str,
    llm_result: Optional[dict] = None,
) -> str:
    """生成完整的 HTML 审核报告（自包含，可直接在浏览器打开）"""
    alert_count = diff_result.alert_count
    if alert_count > 0:
        banner = (
            f'<div class="alert-banner">🚨 发现 {alert_count} 处需关注的差异'
            f'（修改 {diff_result.tampered_count} 处 / 缺失 {diff_result.removed_count} 处 / 新增 {diff_result.added_count} 处）</div>'
        )
    else:
        banner = '<div class="ok-banner">✅ 未发现模版内容被修改，仅有正常填写字段</div>'

    return _HTML_TEMPLATE.format(
        template_name=_escape(template_name),
        contract_name=_escape(contract_name),
        timestamp=datetime.now().strftime("%Y-%m-%d %H:%M"),
        similarity=diff_result.similarity,
        banner=banner,
        filled=diff_result.filled_count,
        tampered=diff_result.tampered_count,
        removed=diff_result.removed_count,
        added=diff_result.added_count,
        sections_html=_build_sections_html(diff_result),
        llm_html=_build_llm_html(llm_result),
    )


# ─────────────────────────────────────────────
# Word 报告
# ─────────────────────────────────────────────
def generate_word_report(
    diff_result: DiffResult,
    template_name: str,
    contract_name: str,
    output_path: str,
    llm_result: Optional[dict] = None,
) -> None:
    """
    生成 Word (.docx) 审核报告，供法务向业务方反馈使用。
    """
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    # ── 页面边距 ──
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    def _add_heading(text, level=1):
        p = doc.add_heading(text, level=level)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _add_colored_para(text, rgb, bold=False, size=11):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.color.rgb = RGBColor(*rgb)
        run.bold = bold
        run.font.size = Pt(size)
        return p

    def _shade_row(row, hex_color: str):
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), hex_color)
            tcPr.append(shd)

    # ── 标题 ──
    _add_heading("合同审核报告", level=1)
    doc.add_paragraph(
        f"模版：{template_name}　　回签合同：{contract_name}\n"
        f"审核时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}　　"
        f"模版相似度：{diff_result.similarity:.1%}"
    ).runs[0].font.color.rgb = RGBColor(107, 114, 128)

    doc.add_paragraph()

    # ── 汇总 ──
    _add_heading("一、审核汇总", level=2)
    tbl = doc.add_table(rows=2, cols=4)
    tbl.style = "Table Grid"
    headers = ["正常填写字段", "疑似修改条款", "内容缺失", "模版外新增"]
    values = [
        str(diff_result.filled_count),
        str(diff_result.tampered_count),
        str(diff_result.removed_count),
        str(diff_result.added_count),
    ]
    colors_h = ["D1FAE5", "FEE2E2", "FEF9C3", "DBEAFE"]
    for i, (h, v, c) in enumerate(zip(headers, values, colors_h)):
        tbl.cell(0, i).text = h
        tbl.cell(1, i).text = v
        # 对每个表头单元格单独着色
        tc = tbl.cell(0, i)._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), c)
        tcPr.append(shd)

    for row in tbl.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)
                    run.bold = True

    doc.add_paragraph()

    # ── 修改条款 ──
    tampered = [i for i in diff_result.items if i.type == DiffType.TAMPERED]
    if tampered:
        _add_heading("二、疑似修改条款（需重点核查）", level=2)
        for idx, item in enumerate(tampered, 1):
            doc.add_paragraph(f"【{idx}】", style="List Number")
            t = doc.add_table(rows=2, cols=2)
            t.style = "Table Grid"
            t.cell(0, 0).text = "模版原文"
            t.cell(0, 1).text = "回签版本"
            t.cell(1, 0).text = item.template_text or ""
            t.cell(1, 1).text = item.contract_text or ""
            _shade_row(t.rows[0], "FEE2E2")
            # 标红修改列
            for run in t.cell(1, 1).paragraphs[0].runs:
                run.font.color.rgb = RGBColor(220, 38, 38)
            doc.add_paragraph()
    else:
        _add_heading("二、修改情况", level=2)
        _add_colored_para("✅ 未发现模版条款被修改。", (22, 163, 74), bold=True)
        doc.add_paragraph()

    # ── 缺失内容 ──
    removed = [i for i in diff_result.items if i.type == DiffType.REMOVED]
    if removed:
        _add_heading("三、模版内容缺失", level=2)
        for item in removed:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(item.template_text or "")
            run.font.color.rgb = RGBColor(180, 83, 9)
        doc.add_paragraph()

    # ── 模版外新增 ──
    added = [i for i in diff_result.items if i.type == DiffType.ADDED]
    if added:
        _add_heading("四、模版外新增内容", level=2)
        for item in added:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(item.contract_text or "")
            run.font.color.rgb = RGBColor(37, 99, 235)
        doc.add_paragraph()

    # ── 填写字段清单 ──
    filled = [i for i in diff_result.items if i.type == DiffType.FILLED]
    # 根据前面已有的章节数动态生成序号
    _section = 3 + (1 if removed else 0) + (1 if added else 0)
    _section_cn = ["", "一", "二", "三", "四", "五", "六"][min(_section, 6)]
    _add_heading(f"{_section_cn}、填写字段清单", level=2)
    if filled:
        tbl2 = doc.add_table(rows=1 + len(filled), cols=2)
        tbl2.style = "Table Grid"
        tbl2.cell(0, 0).text = "模版原文（含空白）"
        tbl2.cell(0, 1).text = "回签填写内容"
        _shade_row(tbl2.rows[0], "D1FAE5")
        for k, item in enumerate(filled, 1):
            tbl2.cell(k, 0).text = item.template_text or ""
            tbl2.cell(k, 1).text = item.contract_text or ""
    else:
        doc.add_paragraph("未检测到填写字段。")

    doc.add_paragraph()

    # ── LLM 结果（可选）──
    if llm_result and "error" not in llm_result:
        _add_heading("附：LLM 辅助分析", level=2)
        summary = llm_result.get("summary", "")
        risk = llm_result.get("risk_level", "")
        rec = llm_result.get("recommendation", "")
        if summary:
            doc.add_paragraph(f"摘要：{summary}")
        if risk:
            risk_label = {"high": "高风险", "medium": "中等风险", "low": "低风险"}.get(risk, risk)
            doc.add_paragraph(f"风险等级：{risk_label}")
        if rec:
            doc.add_paragraph(f"建议：{rec}")

        llm_filled = llm_result.get("filled_fields", [])
        if llm_filled:
            doc.add_paragraph("填写字段（LLM 识别）：")
            tbl3 = doc.add_table(rows=1 + len(llm_filled), cols=3)
            tbl3.style = "Table Grid"
            tbl3.cell(0, 0).text = "字段"
            tbl3.cell(0, 1).text = "填写值"
            tbl3.cell(0, 2).text = "位置"
            _shade_row(tbl3.rows[0], "DBEAFE")
            for k, f in enumerate(llm_filled, 1):
                tbl3.cell(k, 0).text = f.get("field_name", "")
                tbl3.cell(k, 1).text = f.get("value", "")
                tbl3.cell(k, 2).text = f.get("location", "")

    doc.save(output_path)
