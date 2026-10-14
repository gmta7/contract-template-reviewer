import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from modules.date_checker import DateCheckResult
from modules.differ import DiffResult
from modules.llm_analyzer import analyze_with_llm
from modules.reporter import _build_llm_html, generate_word_report
from modules.review_pipeline import format_comment


RISK_RESULT = {
    "contract_type": "服务合同",
    "our_party": "甲方",
    "counterparty": "乙方",
    "transaction_summary": "乙方向甲方提供服务",
    "risk_level": "high",
    "summary": "付款条件存在重大风险。",
    "risk_points": [
        {
            "id": 1,
            "category": "付款",
            "issue_type": "表述不清",
            "title": "付款期限不明确",
            "location": "第四条",
            "clause_excerpt": "甲方应及时付款。",
            "risk_level": "high",
            "reason": "“及时”没有明确期限。",
            "potential_consequence": "可能引发付款逾期争议。",
            "recommendation": "约定明确的付款期限。",
            "suggested_revision": "甲方应在收到合规发票后十个工作日内付款。",
        }
    ],
    "internal_conflicts": [
        {
            "locations": ["第三条", "第四条"],
            "description": "付款时间前后不一致",
            "recommendation": "统一为十个工作日",
        }
    ],
    "items_to_confirm": [
        {"item": "确认服务验收负责人", "reason": "合同未载明"}
    ],
    "recommendation": "修改后签署。",
}


class _FakeCompletions:
    def __init__(self, captured):
        self.captured = captured

    def create(self, **kwargs):
        self.captured.append(kwargs)
        message = SimpleNamespace(content=json.dumps(RISK_RESULT, ensure_ascii=False))
        return SimpleNamespace(choices=[SimpleNamespace(message=message)])


class _FakeClient:
    def __init__(self, captured):
        self.chat = SimpleNamespace(completions=_FakeCompletions(captured))


class LlmReviewModeTests(unittest.TestCase):
    def test_template_and_non_template_use_different_prompts(self):
        captured = []
        fake_client = _FakeClient(captured)
        with patch(
            "modules.openai_client.create_openai_client",
            return_value=fake_client,
        ):
            analyze_with_llm("模版正文", "回签正文", api_key="test", review_mode="template")
            analyze_with_llm("", "普通合同正文", api_key="test", review_mode="non_template")

        template_messages = captured[0]["messages"]
        non_template_messages = captured[1]["messages"]
        self.assertIn("标准合同模版的回签版本", template_messages[0]["content"])
        self.assertIn("【合同模版】", template_messages[1]["content"])
        self.assertIn("没有标准模版可供比对", non_template_messages[0]["content"])
        self.assertIn("【合同全文】", non_template_messages[1]["content"])
        self.assertNotIn("【合同模版】", non_template_messages[1]["content"])

    def test_dingtalk_comment_contains_required_risk_fields(self):
        comment = format_comment(
            contract_name="测试合同.docx",
            is_template=False,
            template_name="",
            similarity=0.0,
            diff_result=DiffResult(),
            meta={"party_a": "甲方公司", "party_b": "乙方公司"},
            date_result=DateCheckResult(info=["日期核查正常"]),
            company_result={"establishment_date": "2020-01-01", "source": "测试"},
            llm_result=RISK_RESULT,
            max_comment_length=10000,
        )
        for expected in (
            "付款期限不明确",
            "位置：第四条",
            "原文：甲方应及时付款。",
            "原因：“及时”没有明确期限。",
            "建议：约定明确的付款期限。",
        ):
            self.assertIn(expected, comment)

    def test_html_and_word_keep_full_risk_details(self):
        html = _build_llm_html(RISK_RESULT)
        for expected in (
            "付款期限不明确",
            "甲方应及时付款。",
            "可能引发付款逾期争议。",
            "甲方应在收到合规发票后十个工作日内付款。",
            "付款时间前后不一致",
            "确认服务验收负责人",
        ):
            self.assertIn(expected, html)

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "report.docx"
            generate_word_report(
                diff_result=DiffResult(),
                template_name="（无匹配模版，AI全文初审）",
                contract_name="测试合同.docx",
                output_path=str(output),
                llm_result=RISK_RESULT,
            )
            doc = Document(output)
            text_parts = [paragraph.text for paragraph in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    text_parts.extend(cell.text for cell in row.cells)
            report_text = "\n".join(text_parts)

        for expected in (
            "付款期限不明确",
            "甲方应及时付款。",
            "可能引发付款逾期争议。",
            "约定明确的付款期限。",
            "甲方应在收到合规发票后十个工作日内付款。",
            "付款时间前后不一致",
            "确认服务验收负责人",
        ):
            self.assertIn(expected, report_text)


if __name__ == "__main__":
    unittest.main()
